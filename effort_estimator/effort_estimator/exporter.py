"""
exporter.py
-----------
Writes the estimate to an Excel workbook (needs:  pip install openpyxl).

Sheet 1 "Estimate"        -> the required 5-column format (one row)
Sheet 2 "Effort Breakdown" -> activity-wise working (formulas, so it stays traceable)
"""

from datetime import date
from typing import List

from calculator import EstimateResult
from config import COMPANY_NAME, EXCEL_HEADERS

PURPLE = "4F2D7F"
PURPLE_LIGHT = "EDE7F6"
BORDER_GREY = "BFBFC8"
FONT = "Arial"


class ExportError(RuntimeError):
    pass


def export_excel(result: EstimateResult, entity_name: str, delivery_date: date,
                 resources: List[str], assumptions: str, path: str) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError as exc:
        raise ExportError("Excel export needs the 'openpyxl' package.\n\n"
                          "Install it once from the VS Code terminal:\n"
                          "    pip install openpyxl") from exc

    thin = Side(style="thin", color=BORDER_GREY)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    head_fill = PatternFill("solid", fgColor=PURPLE)
    body_font = Font(name=FONT, size=10)
    wrap_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrap_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    wb = Workbook()

    # ------------------------------------------------ Sheet 2: breakdown
    bd = wb.active
    bd.title = "Effort Breakdown"
    ws = wb.create_sheet("Estimate", 0)          # Estimate becomes first sheet

    bd["A1"] = f"{COMPANY_NAME} - {result.title}"
    bd["A1"].font = Font(name=FONT, bold=True, size=12, color=PURPLE)
    bd["A2"] = f"Entity: {entity_name}"
    bd["A2"].font = body_font

    row = 4
    for label, value in ([("Mode of Delivery", result.mode)] if result.mode else []) + \
            [(s.label, s.band + ("" if s.direct else f"  (multiplier x{s.multiplier})"))
             for s in result.scopes]:
        bd.cell(row=row, column=1, value=label).font = Font(name=FONT, bold=True, size=10)
        bd.cell(row=row, column=2, value=value).font = body_font
        row += 1

    row += 1
    head_row = row
    for col, text in enumerate(["Activity", "Selected", "Base Effort (hrs)",
                                "Multiplier", "Final Effort (hrs)"], start=1):
        cell = bd.cell(row=head_row, column=col, value=text)
        cell.font, cell.fill, cell.alignment, cell.border = head_font, head_fill, wrap_center, border

    first = head_row + 1
    for i, item in enumerate(result.items):
        r = first + i
        values = [item.particular, item.level, item.base_hours, item.multiplier, f"=C{r}*D{r}"]
        for col, value in enumerate(values, start=1):
            cell = bd.cell(row=r, column=col, value=value)
            cell.font, cell.border = body_font, border
            cell.alignment = wrap_left if col <= 2 else wrap_center
    last = first + len(result.items) - 1

    total_row = last + 1
    days_row = last + 2
    bd.cell(row=total_row, column=1, value="Total Time (in hours)")
    bd.cell(row=total_row, column=5, value=f"=SUM(E{first}:E{last})")
    bd.cell(row=days_row, column=1, value=f"Total Time (in days) @ {result.hours_per_day} hrs/day")
    bd.cell(row=days_row, column=5, value=f"=ROUND(E{total_row}/{result.hours_per_day},2)")
    for r in (total_row, days_row):
        for col in range(1, 6):
            cell = bd.cell(row=r, column=col)
            cell.font = Font(name=FONT, bold=True, size=10)
            cell.fill = PatternFill("solid", fgColor=PURPLE_LIGHT)
            cell.border = border
            cell.alignment = wrap_center if col == 5 else wrap_left

    for col, width in zip("ABCDE", (46, 16, 18, 12, 18)):
        bd.column_dimensions[col].width = width

    # ------------------------------------------------ Sheet 1: required format
    for col, text in enumerate(EXCEL_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font, cell.fill, cell.alignment, cell.border = head_font, head_fill, wrap_center, border

    hours_is_int = float(result.total_hours).is_integer()
    data = [
        entity_name,
        f"='Effort Breakdown'!E{total_row}",       # linked to breakdown total
        delivery_date,
        ", ".join(resources),
        assumptions,
    ]
    for col, value in enumerate(data, start=1):
        cell = ws.cell(row=2, column=col, value=value)
        cell.font, cell.border = body_font, border
        cell.alignment = wrap_left if col in (1, 4, 5) else wrap_center
    ws["B2"].number_format = "00" if hours_is_int else "00.00"
    ws["C2"].number_format = "DD-MMM-YYYY"

    for col, width in zip("ABCDE", (28, 16, 16, 30, 70)):
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 60
    ws.freeze_panes = "A2"

    wb.calculation.fullCalcOnLoad = True          # Excel computes formulas on open
    wb.save(path)
